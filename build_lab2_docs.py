"""Build editable DOCX deliverables from the canonical lab-02 Markdown files.

Markdown and Mermaid sources remain the authoritative project files.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image


ROOT = Path(__file__).resolve().parent / 'docs' / 'lab-02'
BLUE = RGBColor(35, 111, 184)
DARK = RGBColor(30, 51, 73)
MUTED = RGBColor(94, 111, 130)


def clean(text: str) -> str:
    text = re.sub(r'!\[([^]]+)\]\([^)]+\)', r'\1', text)
    text = re.sub(r'\[([^]]+)\]\(([^)]+)\)', r'\1 (\2)', text)
    text = re.sub(r'\*\*([^*]+)\*\*', r'\1', text)
    text = re.sub(r'`([^`]+)`', r'\1', text)
    return text.replace('  ', ' ').strip()


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)


def add_page_field(paragraph) -> None:
    paragraph.add_run('Стр. ')
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    paragraph._p.append(fld)


def configure(doc: Document, short_title: str) -> None:
    sec = doc.sections[0]
    sec.page_height = Cm(29.7)
    sec.page_width = Cm(21)
    sec.top_margin = Cm(1.7)
    sec.bottom_margin = Cm(1.7)
    sec.left_margin = Cm(1.7)
    sec.right_margin = Cm(1.7)
    sec.header_distance = Cm(.8)
    sec.footer_distance = Cm(.8)

    normal = doc.styles['Normal']
    normal.font.name = 'Liberation Sans'
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = DARK
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12
    for style_name, size in [('Title', 19), ('Heading 1', 14), ('Heading 2', 11.5)]:
        style = doc.styles[style_name]
        style.font.name = 'Liberation Sans'
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLUE
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True

    head = sec.header.paragraphs[0]
    head.text = f'GoCue · ЛР2 · {short_title}'
    head.style = normal
    head.runs[0].font.size = Pt(8)
    head.runs[0].font.color.rgb = MUTED
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_field(footer)
    for run in footer.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = MUTED


def add_table(doc: Document, rows: list[list[str]]) -> None:
    columns = max(len(row) for row in rows)
    table = doc.add_table(rows=0, cols=columns)
    table.style = 'Table Grid'
    table.autofit = False
    header = rows[0]
    if columns == 4 and header[:2] == ['ID', 'Приоритет']:
        widths = [1.2, 1.7, 12.4, 2.2]
    elif columns == 3 and header[:2] == ['ID', 'Приоритет']:
        widths = [1.2, 1.7, 14.6]
    elif columns == 4 and header[:2] == ['ID', 'Дано']:
        widths = [1.2, 4.1, 5.1, 7.1]
    elif columns == 4 and header[:2] == ['UC', 'Сценарий']:
        widths = [1.2, 4.3, 5.0, 7.0]
    elif columns == 2:
        widths = [4.3, 13.2]
    else:
        widths = [17.5 / columns] * columns
    for col_i, width in enumerate(widths):
        table.columns[col_i].width = Cm(width)
    for row_i, data in enumerate(rows):
        row = table.add_row()
        tr_pr = row._tr.get_or_add_trPr()
        no_split = OxmlElement('w:cantSplit')
        tr_pr.append(no_split)
        for col_i in range(columns):
            cell = row.cells[col_i]
            cell.width = Cm(widths[col_i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = clean(data[col_i]) if col_i < len(data) else ''
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(1)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    run.font.size = Pt(8)
                    if row_i == 0:
                        run.font.bold = True
                        run.font.color.rgb = BLUE
            if row_i == 0:
                shade(cell, 'EAF2F9')
        if row_i == 0:
            repeat = OxmlElement('w:tblHeader')
            repeat.set(qn('w:val'), 'true')
            tr_pr.append(repeat)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_diagram(doc: Document, image_path: Path, caption: str) -> None:
    png = image_path.with_suffix('.png')
    if not png.exists():
        raise FileNotFoundError(png)
    with Image.open(png) as img:
        width, height = img.size
    max_w, max_h = 17.4, 17.5
    ratio = min(max_w / width, max_h / height)
    figure = doc.add_paragraph()
    figure.alignment = WD_ALIGN_PARAGRAPH.CENTER
    figure.paragraph_format.space_before = Pt(5)
    figure.paragraph_format.space_after = Pt(3)
    figure.add_run().add_picture(str(png), width=Cm(width * ratio), height=Cm(height * ratio))
    fig_caption = doc.add_paragraph(clean(caption))
    fig_caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in fig_caption.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = MUTED


def build(source: Path, target: Path, short_title: str) -> None:
    doc = Document()
    configure(doc, short_title)
    lines = source.read_text(encoding='utf-8').splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith('!['):
            match = re.match(r'!\[([^]]+)\]\(([^)]+)\)', line)
            if match:
                add_diagram(doc, source.parent / match.group(2), match.group(1))
            i += 1
            continue
        if line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                cells = [part.strip() for part in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?', cell) for cell in cells):
                    rows.append(cells)
                i += 1
            if rows:
                add_table(doc, rows)
            continue
        if line.startswith('# '):
            p = doc.add_paragraph(clean(line[2:]), 'Title')
            p.paragraph_format.space_after = Pt(12)
        elif line.startswith('## '):
            if short_title == 'Технический проект' and re.match(r'^## (?:[2-9]|10)\. ', line):
                doc.add_page_break()
            doc.add_paragraph(clean(line[3:]), 'Heading 1')
        elif line.startswith('### '):
            doc.add_paragraph(clean(line[4:]), 'Heading 2')
        elif line.startswith('- '):
            doc.add_paragraph(clean(line[2:]), 'List Bullet')
        elif re.match(r'^\d+\. ', line):
            doc.add_paragraph(clean(re.sub(r'^\d+\. ', '', line)), 'List Number')
        elif line.startswith('**Лабораторная работа'):
            p = doc.add_paragraph(clean(line))
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = MUTED
        else:
            doc.add_paragraph(clean(line))
        i += 1
    doc.save(target)
    print(target)


if __name__ == '__main__':
    build(ROOT / 'Техническое_задание.md', ROOT / 'Техническое_задание.docx', 'Техническое задание')
    build(ROOT / 'Технический_проект.md', ROOT / 'Технический_проект.docx', 'Технический проект')
