"""Update the existing GoCue LR1 report from verified repository/project facts.

Unlike build_lab1_report.py, this script opens the current DOCX and preserves
manual edits outside the targeted cells and sections. Run from any directory.
"""

from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parent
REPORT = ROOT / "docs" / "lab-01" / "ЛР1_Отчёт_GoCue.docx"
document = Document(REPORT)
tables = document.tables


def set_row(table_index, row_index, values):
    for cell, value in zip(tables[table_index].rows[row_index].cells, values):
        cell.text = str(value)


def paragraph(text):
    return next(p for p in document.paragraphs if p.text == text)


# The logins were explicitly confirmed by the repository owner.
for index, login in enumerate(("BamStalker", "techmo0", "whooaaa1", "iamtsunami-1"), start=1):
    name = tables[17].rows[index].cells[0].text.split(" (")[0]
    tables[17].rows[index].cells[0].text = f"{name} (@{login})"

# Record this use of AI without attributing actions to other team members.
journal = tables[21]
if not any("GitHub Projects и разделов 6–7" in row.cells[1].text for row in journal.rows[1:]):
    cells = journal.add_row().cells
    values = (
        "14.10.2026",
        "GitHub Projects и разделов 6–7; ChatGPT / Codex",
        "Созданы issues, настроена доска, подготовлены правки отчёта",
        "Принято после проверки фактических статусов; действия команды не подменены",
        "Сверено с GitHub Projects, списком issues и методичкой",
    )
    for cell, value in zip(cells, values):
        cell.text = value
        for run in cell.paragraphs[0].runs:
            run.font.size = Pt(8.5)

# 6.1. Links and access. A public repository is not proof of access to a
# private user-owned Project, so the teacher's access remains unconfirmed.
tables[22].rows[0].cells[2].text = "Кто ведёт"
tables[22].rows[0].cells[3].text = "Доступ преподавателя"
set_row(22, 1, (
    "Организация GitHub",
    "Не создавалась; проект размещён в личном аккаунте techmo0",
    "Боровик Ф. И.",
    "—",
))
set_row(22, 2, (
    "Репозиторий",
    "https://github.com/techmo0/GoCue (публичный)",
    "Боровик Ф. И.",
    "☑ доступен по ссылке",
))
set_row(22, 3, (
    "Доска GitHub Projects",
    "https://github.com/users/techmo0/projects/1/views/1 (приватная; права редактора у всей команды)",
    "Боровик Ф. И.",
    "☐ не подтверждён",
))
set_row(22, 4, (
    "Зеркало репозитория",
    "Не настроено; для ЛР 1 необязательно",
    "—",
    "—",
))
set_row(22, 5, (
    "Чат команды",
    "Telegram; ссылка не публикуется",
    "Бартоломей А. М.",
    "—",
))

# 6.2. README contents must describe the repository, not a future plan.
tables[23].rows[0].cells[0].text = (
    "README.md объясняет проблему забытых поручений, идею GoCue, состав команды с GitHub-логинами, "
    "границы MVP и выбранную платформу Android. Он ссылается на спецификацию, отчёт ЛР 1, "
    "перечень issues, основную доску и табличный бэклог. В README отдельно отмечены "
    "неподтверждённые действия: рецензия, ретроспектива, доступ преподавателя и сдача на портал."
)

# 6.3. Verified project settings.
settings = [
    ("☑", "Связан с techmo0/GoCue."),
    ("☑", "Пять статусов; доска сгруппирована по Status."),
    ("☑", "Высокий, Средний, Низкий."),
    ("☑", "ЛР 1: 12–25.10; ЛР 2: 26.10–08.11.2026. Это план, не сроки курса."),
    ("☑", "Числовая оценка в часах."),
    ("☑", "Вид № 1; ссылка в п. 6.1."),
    ("☑", "Вид № 2; сортировка по Priority."),
    ("☑", "Workflow Item closed включён."),
    ("☑", "#1–#14 — ЛР 1; #15–#18 — прототип."),
    ("☐", "Пока есть коммиты только @techmo0; задача #14 открыта."),
]
for row_index, (done, note) in enumerate(settings, start=1):
    tables[24].rows[row_index].cells[1].text = done
    tables[24].rows[row_index].cells[2].text = note
for column, width in zip(tables[24].columns, (Cm(9.5), Cm(2.5), Cm(5.0))):
    column.width = width
for row in tables[24].rows:
    for cell, width in zip(row.cells, (Cm(9.5), Cm(2.5), Cm(5.0))):
        cell.width = width

# 6.4. Report the actual board state. One issue has one accountable assignee;
# work requiring team approval is described in the issue itself.
next(p for p in document.paragraphs if p.text.startswith((
    "Исполнители указаны предварительно.", "Номера, исполнители, оценки и статусы сверены"
))).text = (
    "Номера, исполнители, оценки и статусы сверены с GitHub Projects 14.10.2026. "
    "Оценки предварительные; Done означает завершение конкретной задачи, но не всей ЛР 1."
)
lab_tasks = [
    (1, "Подтвердить матрицу компетенций и роли", "Бартоломей А. М.; вся команда", "1", "Ready"),
    (2, "Зафиксировать карточку проекта и тему", "Бартоломей А. М.", "2", "Done"),
    (3, "Подготовить репозиторий и README", "Боровик Ф. И.", "1", "Done"),
    (4, "Настроить доску и доступ команды", "Боровик Ф. И.", "2", "Done"),
    (5, "Согласовать Vision", "Бартоломей А. М.", "3", "In review"),
    (6, "Получить рецензию Vision другой команды", "Золотилина Е. С.", "1", "Backlog"),
    (7, "Согласовать устав проекта", "Бартоломей А. М.", "2,5", "In review"),
    (8, "Утвердить соглашения и критерии готовности", "Бартоломей А. М.; вся команда", "1,5", "Ready"),
    (9, "Утвердить политику ИИ и журнал", "Поспелова М. Т.", "1", "In review"),
    (10, "Проверить отчёт и загрузить на портал", "Боровик Ф. И.", "2", "In progress"),
    (11, "Провести ретроспективу", "Бартоломей А. М.; вся команда", "0,5", "Backlog"),
    (12, "Проверить доступ преподавателя к доске", "Боровик Ф. И.", "0,5", "Backlog"),
    (13, "Проверить сценарии с добровольцами", "Бартоломей А. М.", "2", "Backlog"),
    (14, "Получить коммит от каждого участника", "Бартоломей А. М.; вся команда", "1", "Backlog"),
]
task_table = tables[25]
task_table.rows[0].cells[0].text = "№"
while len(task_table.rows) < len(lab_tasks) + 1:
    task_table.add_row()
for row_index, (number, title, owner, hours, status) in enumerate(lab_tasks, start=1):
    for cell, value in zip(task_table.rows[row_index].cells, (f"#{number}", title, owner, hours, status)):
        cell.text = value

# 6.5. An API check is not passed off as a browser screenshot.
tables[26].rows[0].cells[0].text = (
    "Доска: https://github.com/users/techmo0/projects/1/views/1; "
    "табличный бэклог: https://github.com/users/techmo0/projects/1/views/2. "
    "Проверено через GitHub API: 18 карточек, из них 14 для ЛР 1 и 4 для прототипа; "
    "все назначены, имеют статус, приоритет, оценку и итерацию. "
    "Скриншот интерфейса ещё нужно сделать в браузере под аккаунтом с доступом к приватной доске "
    "и вставить сюда перед сдачей; API-сводка не выдаётся за скриншот."
)

# 7. Keep the retrospect and individual contributions factual. The team's
# meeting and portal submission cannot be completed by editing a document.
intro = next(p for p in document.paragraphs if p.text.startswith((
    "К моменту подготовки отчёта тема GoCue согласована", "Состояние на 14.10.2026: тема согласована"
)))
intro.text = (
    "Состояние на 14.10.2026: тема согласована, выбрана платформа Android, создан публичный "
    "репозиторий и настроена приватная доска с задачами. Отчёт подготовлен к проверке командой. "
    "Рецензия другой команды, мини-защита, ретроспектива и загрузка на портал ещё не подтверждены."
)
section_8 = paragraph("8. Источники для раздела об аналогах")
subheading_style = paragraph("6.1. Ссылки и доступы").style


def before_sources(element):
    section_8._p.addprevious(element)


def add_text(text, style=None):
    new = document.add_paragraph(text, style=style)
    before_sources(new._p)
    for run in new.runs:
        if style is None:
            run.font.size = Pt(10)
    return new


def add_table(headers, rows, widths):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = task_table.style
    table.autofit = False
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
    for values in rows:
        for cell, value in zip(table.add_row().cells, values):
            cell.text = value
    for column, width in zip(table.columns, widths):
        column.width = width
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = width
            for paragraph_ in cell.paragraphs:
                for run in paragraph_.runs:
                    run.font.size = Pt(8.5)
    before_sources(table._element)


if not any(p.text.startswith("7.1.") for p in document.paragraphs):
    add_text("7.1. Итоги итерации", subheading_style)
    add_text(
        "Выполнено: зафиксированы тема и карточка проекта (#2), репозиторий и README (#3), "
        "доска с доступом команды (#4). Подготовлены черновики Vision и устава; они находятся "
        "на проверке (#5, #7). Матрица компетенций заполнена, но ждёт подтверждения всех участников (#1)."
    )
    add_text(
        "Ретроспектива «что получилось / что мешало / что меняем» будет записана командой "
        "после занятия в задаче #11 и перенесена сюда. До её проведения нельзя считать эти выводы "
        "общим решением команды."
    )

    add_text("7.2. Вклад участников", subheading_style)
    add_text(
        "Ниже отделён подтверждённый результат от назначенных задач. Оценки Estimate на доске — "
        "плановые часы, а не фактически отработанное время."
    )
    add_table(
        ("Участник", "Подтверждённый вклад", "Задачи на доске", "Коммиты"),
        (
            ("Боровик Ф. И.", "Репозиторий, README, настройка доски; правка отчёта", "#3, #4, #10, #12", "Есть в истории репозитория"),
            ("Бартоломей А. М.", "Личный вклад требует подтверждения участником", "#1, #2, #5, #7, #8, #11, #13, #14", "Пока не выявлены"),
            ("Золотилина Е. С.", "Личный вклад требует подтверждения участником", "#6; прототип #17", "Пока не выявлены"),
            ("Поспелова М. Т.", "Личный вклад требует подтверждения участником", "#9; прототип #16", "Пока не выявлены"),
        ),
        (Cm(3.0), Cm(5.0), Cm(4.5), Cm(4.0)),
    )

    add_text("7.3. Проверка перед сдачей", subheading_style)
    add_table(
        ("Проверка", "Состояние на 14.10.2026", "Следующее действие"),
        (
            ("Тема и проектные документы", "Тема согласована; Vision и устав на проверке", "Команде утвердить #1, #5, #7–#9"),
            ("Репозиторий и доска", "Ссылки работают; 18 задач, доступ команды настроен", "Проверить доступ преподавателя (#12)"),
            ("Коммиты участников", "Подтверждены только коммиты @techmo0", "Каждому сделать собственный коммит (#14)"),
            ("Рецензия и ретроспектива", "Не проведены", "Выполнить #6 и #11 на занятии"),
            ("Скриншот доски", "Не приложен", "Вставить снимок реальной доски в п. 6.5"),
            ("Отчёт и журнал ИИ", "DOCX/PDF подготовлены; журнал дополнен", "Команде вычитать и подтвердить"),
            ("Мини-защита и портал", "Не подтверждены", "Провести защиту и загрузить отчёт (#10)"),
        ),
        (Cm(4.0), Cm(6.0), Cm(6.5)),
    )

document.core_properties.title = "ЛР 1. Инициация проекта и команды — GoCue"
document.core_properties.subject = "Проектный практикум, команда Диалог"
document.core_properties.author = "Команда «Диалог»"
document.save(REPORT)
print(REPORT)
